"""Verify P14 milestone provenance, numbers, text and layout; produce previews."""
import argparse
import hashlib
import json
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from review_research_presentation import normalized

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / 'materials/20261021_p14'
STEM = 'AI_harness_P11_P14_20261021'


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--preview-dir', type=Path, required=True)
    args = parser.parse_args()
    manifest = json.loads((FOLDER/'build_manifest.json').read_text())
    specs = json.loads((FOLDER/'slides.json').read_text())
    pdf = pymupdf.open(FOLDER/(STEM+'.pdf'))
    pptx = Presentation(FOLDER/(STEM+'.pptx'))
    assert len(pdf) == len(specs) == len(pptx.slides) == manifest['pages'] == 11
    for field in ('source_sha256', 'builder_source_sha256'):
        for path, expected in manifest[field].items():
            assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest() == expected, path
    for name, expected in manifest['figures_sha256'].items():
        assert hashlib.sha256((FOLDER/'figures'/(name+'.png')).read_bytes()).hexdigest() == expected
    table_text = normalized(pdf[3].get_text())
    for phase, row in zip((11,12,13,14), manifest['totals'], strict=True):
        data = json.loads((ROOT/f'reports/P{phase}_final_measurements.json').read_text())
        assert not data['pending'] and not data['audit_errors']
        assert row['tokens'] == sum(t['known_total_tokens'] for t in data['trials'])
        assert row['unknown'] == sum(t['unknown_usage_calls'] for t in data['trials'])
        assert row['successes'] == sum(t['evaluation']['official_success'] for t in data['trials'])
        shown = ('≥' if row['unknown'] else '') + f"{row['tokens']:,}"
        assert normalized(shown) in table_text, shown
    interrupted = json.loads((ROOT/'reports/P13_interruption_20261021.json').read_text())
    known = interrupted['interrupted_cost']['known_total_tokens']
    assert normalized(f'{known:,}') in normalized(pdf[4].get_text())
    overlaps = []
    for n, (spec, page, slide) in enumerate(zip(specs,pdf,pptx.slides),1):
        pt = normalized(page.get_text())
        st = normalized(''.join(s.text for s in slide.shapes if s.has_text_frame))
        items = [e for e in spec['elements'] if e['kind']=='text']
        for a in items:
            assert normalized(a['text']) in pt, (n,'pdf',a['text'])
            assert normalized(a['text']) in st, (n,'pptx',a['text'])
        for i, a in enumerate(items):
            ah = len(a['rendered_lines'])*a['size']*a.get('leading',1.32)
            for b in items[i+1:]:
                bh = len(b['rendered_lines'])*b['size']*b.get('leading',1.32)
                dx = min(a['x']+a['w'],b['x']+b['w'])-max(a['x'],b['x'])
                dy = min(a['y']+ah,b['y']+bh)-max(a['y'],b['y'])
                if dx > 1 and dy > 1:
                    overlaps.append({'page':n,'a':a['text'],'b':b['text']})
    assert not overlaps, overlaps
    links = [x['uri'] for page in pdf for x in page.get_links() if 'uri' in x]
    assert len(links)==5
    prefix = 'https://github.com/sophiria1225/AI-/blob/main/'
    for url in links:
        assert url.startswith(prefix) and (ROOT/url.removeprefix(prefix)).is_file(), url
    args.preview_dir.mkdir(parents=True,exist_ok=True)
    for start in range(0,len(pdf),6):
        sheet = Image.new('RGB',(1200,1050),'#D8E0E8');draw=ImageDraw.Draw(sheet)
        for k,page in enumerate(pdf[start:start+6]):
            pix=page.get_pixmap(matrix=pymupdf.Matrix(.6,.6),alpha=False)
            im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
            x,y=(k%2)*600+12,(k//2)*350+24
            sheet.paste(im,(x,y));draw.text((x,y-18),str(start+k+1),fill='black')
        sheet.save(args.preview_dir/f'contact_{start+1:02d}.png')
    for n,page in enumerate(pdf,1):
        page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False).save(args.preview_dir/f'slide_{n:02d}.png')
    report={'pages':len(pdf),'source_hashes_match':True,'numeric_checks_passed':True,
            'unknown_costs_marked_as_lower_bounds':True,'interrupted_cost_included':True,
            'all_pdf_and_pptx_text_present':True,'text_overlaps':overlaps,
            'valid_source_links':len(links),'native_powerpoint_rendered':False,
            'limitation':'PDF previews require visual review; native PowerPoint fonts may vary.',
            'artifact_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in FOLDER.glob(STEM+'.*')}}
    (FOLDER/'review.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report))


if __name__=='__main__':
    main()
