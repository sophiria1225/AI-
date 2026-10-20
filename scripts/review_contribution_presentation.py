"""Validate the approach deck and produce PDF previews for human visual review."""
import argparse
import csv
import hashlib
import json
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from review_research_presentation import normalized

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "materials/20261020_revised"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--preview-dir", type=Path, required=True)
    args = parser.parse_args()
    data = json.loads((FOLDER / "data_sources.json").read_text())
    specs = json.loads((FOLDER / "slides.json").read_text())
    pdf = pymupdf.open(FOLDER / "AI_harness_research_20261020.pdf")
    pptx = Presentation(FOLDER / "AI_harness_research_20261020.pptx")
    assert len(pdf) == len(specs) == len(pptx.slides) == 16
    for source, sha in data["source_sha256"].items():
        assert hashlib.sha256((ROOT/source).read_bytes()).hexdigest() == sha, source
    with (FOLDER / "observations.csv").open() as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == len(data["observations"]) == 6
    for csv_row, row in zip(rows, data["observations"]):
        assert csv_row["task"] == row["task"]
        assert int(csv_row["total_tokens"]) == row["total_tokens"]
        assert csv_row["resolved"] == str(row["resolved"])
    # Source files are also hash checked. Validate the plotted observations.
    p3 = json.loads((ROOT / "reports/P3_measurements.json").read_text())
    p5 = json.loads((ROOT / "reports/P5_pro_measurements.json").read_text())
    originals = [(r["known_total_tokens"], r["evaluation"]["resolved"]) for r in p3["runs"]]
    originals += [(r["result"]["measurements"]["known_total_tokens"],
                   r["result"]["measurements"]["evaluation"]["resolved"]) for r in p5["runs"]]
    assert originals == [(r["total_tokens"], r["resolved"]) for r in data["observations"]]
    overlaps, checks = [], []
    for n, (spec, page, slide) in enumerate(zip(specs, pdf, pptx.slides), 1):
        pdf_text = normalized(page.get_text())
        pptx_text = normalized("".join(sh.text for sh in slide.shapes if sh.has_text_frame))
        elements = [e for e in spec["elements"] if e["kind"] == "text"]
        for e in elements:
            assert normalized(e["text"]) in pdf_text, (n, "pdf", e["text"])
            assert normalized(e["text"]) in pptx_text, (n, "pptx", e["text"])
        for j, a in enumerate(elements):
            ah = len(a["rendered_lines"]) * a["size"] * a.get("leading", 1.32)
            for b in elements[j+1:]:
                bh = len(b["rendered_lines"]) * b["size"] * b.get("leading", 1.32)
                dx = min(a["x"]+a["w"], b["x"]+b["w"]) - max(a["x"], b["x"])
                dy = min(a["y"]+ah, b["y"]+bh) - max(a["y"], b["y"])
                if dx > 1 and dy > 1:
                    overlaps.append({"page": n, "a": a["text"], "b": b["text"]})
        checks.append({"page": n, "pdf_text_present": True, "pptx_text_present": True})
    assert not overlaps, overlaps
    all_text = normalized("".join(p.get_text() for p in pdf))
    for required in ("580,169", "571,839", "5,050.03", "509.62", "14.58%", "96.05%"):
        assert required in all_text
    assert sum(len(p.get_links()) for p in pdf) == 10
    args.preview_dir.mkdir(parents=True, exist_ok=True)
    for start in range(0, len(pdf), 8):
        sheet = Image.new("RGB", (1200, 1000), "#D8E0E8")
        draw = ImageDraw.Draw(sheet)
        for k, page in enumerate(pdf[start:start+8]):
            pix = page.get_pixmap(matrix=pymupdf.Matrix(.6, .6), alpha=False)
            im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            im.thumbnail((576, 220))
            x, y = (k % 2)*600+12, (k//2)*250+25
            sheet.paste(im, (x, y))
            draw.text((x, y-18), str(start+k+1), fill="black")
        sheet.save(args.preview_dir / f"contact_{start+1:02d}.png")
    for n, page in enumerate(pdf, 1):
        page.get_pixmap(matrix=pymupdf.Matrix(1.25, 1.25), alpha=False).save(
            args.preview_dir / f"slide_{n:02d}.png")
    for page_number, name in [(5, "proposed_architecture"), (6, "four_conditions"),
                              (7, "central_plan_vs_transfer")]:
        pdf[page_number].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).save(
            FOLDER / "figures" / f"{name}.png")
    report = {
        "pages": len(pdf), "checks": checks, "source_hashes_match": True,
        "numeric_checks_passed": True, "text_overlaps": overlaps,
        "source_links": 10, "pdf_previews_created": 16,
        "native_powerpoint_rendered": False,
        "limitation": "Package/text verified. Native PowerPoint display may depend on fonts. Visual review is separate.",
        "artifact_sha256": {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                           for p in FOLDER.glob("AI_harness_research_20261020.*")},
    }
    (FOLDER / "review.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"pages": len(pdf), "text": "present", "overlaps": len(overlaps),
                      "numeric_checks": "passed", "source_links": 10}))


if __name__ == "__main__":
    main()
