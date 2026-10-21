"""Check export milestone sources, all numeric rows and rendered artifacts."""
import argparse
import hashlib
import json
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation

from build_export_presentation import ROOT, OUTPUT, STEM, sources
from review_research_presentation import normalized


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--preview-dir',type=Path,required=True)
    args=parser.parse_args()
    manifest=json.loads((OUTPUT/'build_manifest.json').read_text())
    specs=json.loads((OUTPUT/'slides.json').read_text())
    pdf=pymupdf.open(OUTPUT/(STEM+'.pdf'))
    deck=Presentation(OUTPUT/(STEM+'.pptx'))
    assert len(specs)==len(pdf)==len(deck.slides)==manifest['pages']==6
    for field in ('source_sha256','builder_source_sha256'):
        for path,expected in manifest[field].items():
            assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==expected,path
    for name,expected in manifest['figures_sha256'].items():
        assert hashlib.sha256((OUTPUT/'figures'/(name+'.png')).read_bytes()).hexdigest()==expected
    rows,groups,_=sources()
    assert len(rows)==12
    table_text=normalized(pdf[2].get_text())
    for g in groups:
        for value in (f"{g['cap']:,}",f"{g['submitted']}/3",f"{g['correct']}/3",
                      f"{g['input_tokens']+g['output_tokens']:,}"):
            assert normalized(value) in table_text,value
    for phase in ('P20','P21'):
        total=sum(r['input_tokens']+r['output_tokens'] for r in rows if r['phase']==phase)
        assert normalized(f'{total:,}') in normalized(pdf[0].get_text())
    overlaps=[]
    for n,(spec,page,slide) in enumerate(zip(specs,pdf,deck.slides),1):
        pdf_text=normalized(page.get_text())
        ppt_text=normalized(''.join(s.text for s in slide.shapes if s.has_text_frame))
        items=[e for e in spec['elements'] if e['kind']=='text']
        for a in items:
            assert normalized(a['text']) in pdf_text,(n,'PDF',a['text'])
            assert normalized(a['text']) in ppt_text,(n,'PowerPoint',a['text'])
        for i,a in enumerate(items):
            ah=len(a['rendered_lines'])*a['size']*a.get('leading',1.32)
            for b in items[i+1:]:
                bh=len(b['rendered_lines'])*b['size']*b.get('leading',1.32)
                dx=min(a['x']+a['w'],b['x']+b['w'])-max(a['x'],b['x'])
                dy=min(a['y']+ah,b['y']+bh)-max(a['y'],b['y'])
                if dx>1 and dy>1:
                    overlaps.append({'page':n,'a':a['text'],'b':b['text']})
    assert not overlaps,overlaps
    links=[x['uri'] for page in pdf for x in page.get_links() if 'uri' in x]
    assert len(links)==4
    prefix='https://github.com/sophiria1225/AI-/blob/main/'
    for url in links:
        assert url.startswith(prefix) and (ROOT/url.removeprefix(prefix)).is_file(),url
    args.preview_dir.mkdir(parents=True,exist_ok=True)
    contact=Image.new('RGB',(1200,1050),'#D8E0E8');draw=ImageDraw.Draw(contact)
    for n,page in enumerate(pdf):
        pix=page.get_pixmap(matrix=pymupdf.Matrix(.6,.6),alpha=False)
        image=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
        x,y=(n%2)*600+12,(n//2)*350+24
        contact.paste(image,(x,y));draw.text((x,y-18),str(n+1),fill='black')
        page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False).save(args.preview_dir/f'slide_{n+1:02d}.png')
    contact.save(args.preview_dir/'contact_01.png')
    report={'pages':len(pdf),'source_hashes_match':True,'numeric_checks_passed':True,
            'all_twelve_generations_retained':True,'all_pdf_and_pptx_text_present':True,
            'text_overlaps':overlaps,'valid_source_links':len(links),
            'native_powerpoint_rendered':False,'pdf_visually_reviewed':False,
            'review_script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
            'limitation':'Visual review of PDF previews remains required; no native PowerPoint rendering.',
            'artifact_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in OUTPUT.glob(STEM+'.*')}}
    (OUTPUT/'review.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'pages':len(pdf),'sources':'match','numeric_checks':'passed','overlaps':len(overlaps)}))


if __name__=='__main__':
    main()
