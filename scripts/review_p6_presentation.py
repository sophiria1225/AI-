"""Check the completed-P6 artifacts and prepare private previews for visual review."""
import argparse
import hashlib
import json
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from review_research_presentation import normalized

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / 'materials/20261021_p6'
STEM = 'AI_harness_P6_20261021'


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--preview-dir', type=Path, required=True)
    args = parser.parse_args()
    manifest = json.loads((FOLDER / 'build_manifest.json').read_text())
    specs = json.loads((FOLDER / 'slides.json').read_text())
    pdf = pymupdf.open(FOLDER / (STEM + '.pdf'))
    pptx = Presentation(FOLDER / (STEM + '.pptx'))
    assert len(pdf) == len(specs) == len(pptx.slides) == manifest['pages'] == 11
    for field in ('source_sha256', 'builder_source_sha256'):
        for path, expected in manifest[field].items():
            assert hashlib.sha256((ROOT / path).read_bytes()).hexdigest() == expected, path
    for name, expected in manifest['figures_sha256'].items():
        assert hashlib.sha256((FOLDER / 'figures' / (name + '.png')).read_bytes()).hexdigest() == expected
    results = json.loads((ROOT / 'reports/P6_final_measurements.json').read_text())
    assert not results['pending'] and results['total_model_calls'] == 114
    assert results['known_total_tokens'] == 787520 and results['unknown_usage_calls'] == 4
    result_text = normalized(pdf[4].get_text())
    for row in results['trials']:
        assert normalized(f"{row['known_total_tokens']:,}") in result_text
    overlaps = []
    for n, (spec, page, slide) in enumerate(zip(specs, pdf, pptx.slides), 1):
        pdf_text = normalized(page.get_text())
        pptx_text = normalized(''.join(sh.text for sh in slide.shapes if sh.has_text_frame))
        elements = [e for e in spec['elements'] if e['kind'] == 'text']
        for e in elements:
            assert normalized(e['text']) in pdf_text, (n, 'pdf', e['text'])
            assert normalized(e['text']) in pptx_text, (n, 'pptx', e['text'])
        for index, a in enumerate(elements):
            ah = len(a['rendered_lines']) * a['size'] * a.get('leading', 1.32)
            for b in elements[index+1:]:
                bh = len(b['rendered_lines']) * b['size'] * b.get('leading', 1.32)
                dx = min(a['x']+a['w'], b['x']+b['w']) - max(a['x'], b['x'])
                dy = min(a['y']+ah, b['y']+bh) - max(a['y'], b['y'])
                if dx > 1 and dy > 1:
                    overlaps.append({'page': n, 'a': a['text'], 'b': b['text']})
    assert not overlaps, overlaps
    links = [link['uri'] for page in pdf for link in page.get_links() if 'uri' in link]
    assert len(links) == 5
    prefix = 'https://github.com/sophiria1225/AI-/blob/main/'
    for link in links:
        assert link.startswith(prefix) and (ROOT / link.removeprefix(prefix)).is_file(), link
    args.preview_dir.mkdir(parents=True, exist_ok=True)
    for start in range(0, len(pdf), 6):
        sheet = Image.new('RGB', (1200, 1050), '#D8E0E8')
        draw = ImageDraw.Draw(sheet)
        for k, page in enumerate(pdf[start:start+6]):
            pix = page.get_pixmap(matrix=pymupdf.Matrix(.6, .6), alpha=False)
            im = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
            im.thumbnail((576, 324))
            x, y = (k % 2)*600+12, (k//2)*350+24
            sheet.paste(im, (x, y))
            draw.text((x, y-18), str(start+k+1), fill='black')
        sheet.save(args.preview_dir / f'contact_{start+1:02d}.png')
    for n, page in enumerate(pdf, 1):
        page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).save(
            args.preview_dir / f'slide_{n:02d}.png')
    report = {
        'pages': len(pdf), 'source_hashes_match': True,
        'all_pdf_and_pptx_text_present': True, 'p6_numeric_checks_passed': True,
        'text_overlaps': overlaps, 'valid_local_source_links': len(links),
        'pdf_previews_created': len(pdf), 'native_powerpoint_rendered': False,
        'limitation': 'PDF previews require separate visual review. Native PowerPoint fonts may vary.',
        'artifact_sha256': {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                           for p in FOLDER.glob(STEM + '.*')},
    }
    (FOLDER / 'review.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(report))


if __name__ == '__main__':
    main()
