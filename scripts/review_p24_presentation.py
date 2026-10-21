"""Check source hashes, numbers and layout before visual review of P24 slides."""
import argparse
import hashlib
import json
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from review_research_presentation import normalized
from build_p24_presentation import ROOT, FOLDER, STEM, sources


def main():
    p=argparse.ArgumentParser()
    p.add_argument('--preview-dir',type=Path,required=True)
    args=p.parse_args()
    pair,note,_=sources()
    manifest=json.loads((FOLDER/'build_manifest.json').read_text())
    specs=json.loads((FOLDER/'slides.json').read_text())
    pdf=pymupdf.open(FOLDER/(STEM+'.pdf'))
    deck=Presentation(FOLDER/(STEM+'.pptx'))
    assert len(pdf)==len(specs)==len(deck.slides)==manifest['pages']==8
    for key in ('source_sha256','builder_source_sha256'):
        for path,expected in manifest[key].items():
            assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==expected,path
    for name,expected in manifest['figures_sha256'].items():
        assert hashlib.sha256((FOLDER/'figures'/name).read_bytes()).hexdigest()==expected
    tab=normalized(pdf[3].get_text())
    for label,row in zip(note['task_labels'],pair['pairs'],strict=True):
        assert normalized(label) in tab
        for arm in ('baseline','guided'):
            assert normalized(f"{row[arm]['known_total_tokens']:,}") in tab
    for arm in ('baseline','guided'):
        assert normalized(f"{pair['totals'][arm]['successes']}/6") in tab
        assert normalized(f"{pair['totals'][arm]['known_total_tokens']:,}") in tab
    overlaps=[]
    for n,(spec,page,slide) in enumerate(zip(specs,pdf,deck.slides),1):
        pdf_text=normalized(page.get_text())
        ppt_text=normalized(''.join(s.text for s in slide.shapes if s.has_text_frame))
        items=[e for e in spec['elements'] if e['kind']=='text']
        for a in items:
            assert normalized(a['text']) in pdf_text,(n,'pdf',a['text'])
            assert normalized(a['text']) in ppt_text,(n,'pptx',a['text'])
        for i,a in enumerate(items):
            ah=len(a['rendered_lines'])*a['size']*a.get('leading',1.32)
            for b in items[i+1:]:
                bh=len(b['rendered_lines'])*b['size']*b.get('leading',1.32)
                dx=min(a['x']+a['w'],b['x']+b['w'])-max(a['x'],b['x'])
                dy=min(a['y']+ah,b['y']+bh)-max(a['y'],b['y'])
                if dx>1 and dy>1:
                    overlaps.append({'page':n,'a':a['text'],'b':b['text']})
    assert not overlaps,overlaps
    links=[x['uri'] for page in pdf for x in page.get_links() if 'uri' in x]
    assert len(links)==4
    prefix='https://github.com/sophiria1225/AI-/blob/main/'
    for url in links:
        assert url.startswith(prefix) and (ROOT/url.removeprefix(prefix)).is_file(),url
    args.preview_dir.mkdir(parents=True,exist_ok=True)
    for start in range(0,len(pdf),4):
        sheet=Image.new('RGB',(1200,700),'#D8E0E8');draw=ImageDraw.Draw(sheet)
        for k,page in enumerate(pdf[start:start+4]):
            pix=page.get_pixmap(matrix=pymupdf.Matrix(.6,.6),alpha=False)
            im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
            x,y=(k%2)*600+12,(k//2)*350+24
            sheet.paste(im,(x,y));draw.text((x,y-18),str(start+k+1),fill='black')
        sheet.save(args.preview_dir/f'contact_{start+1:02d}.png')
    for n,page in enumerate(pdf,1):
        page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False).save(args.preview_dir/f'slide_{n:02d}.png')
    report=dict(pages=len(pdf),source_hashes_match=True,numeric_checks_passed=True,
        all_six_pairs_retained=True,all_pdf_and_pptx_text_present=True,text_overlaps=overlaps,
        valid_source_links=len(links),native_powerpoint_rendered=False,pdf_visually_reviewed=False,
        review_script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        limitation='PDF previews need visual review; native PowerPoint font rendering may vary.',
        artifact_sha256={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in FOLDER.glob(STEM+'.*')})
    (FOLDER/'review.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report))


if __name__=='__main__':
    main()
