#!/usr/bin/env python3
"""Build an editable one-column Word companion from the TeX manuscript.

Requires Python packages python-docx and lxml, and an executable Pandoc
(set PANDOC or place pandoc on PATH). Does not change scientific data or figures.
The IEICE PDF is the typeset manuscript; this Word file is an editing companion.
"""
from __future__ import annotations
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import zipfile
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parent
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read(name):
    return (ROOT / name).read_text(encoding='utf-8')


def make_source():
    main = read('main.tex')
    body = main.split(r'\begin{document}', 1)[1].split(r'\end{document}', 1)[0]
    body = body.replace(r'\maketitle', '').replace(r'\input{abstracts.tex}', '')
    body = body.replace(r'\input{tables.tex}', read('tables.tex'))
    refs = re.findall(r'\\bibitem\{([^}]+)\}\n(.*?)(?=\\bibitem|\\end\{thebibliography\})', read('references.tex'), re.S)
    numbers = {key: index + 1 for index, (key, _) in enumerate(refs)}
    body = re.sub(r'\\cite\{([^}]+)\}', lambda m:'['+', '.join(str(numbers[k]) for k in m[1].split(','))+']', body)
    labels = {'fig:architecture':'1', 'fig:results':'2', 'tab:summary':'1', 'tab:termination':'2', 'eq:cost':'1'}
    body = re.sub(r'\\eqref\{([^}]+)\}', lambda m:'('+labels[m[1]]+')', body)
    body = re.sub(r'\\ref\{([^}]+)\}', lambda m:labels[m[1]], body)
    body = re.sub(r'\\label\{[^}]+\}', '', body)
    body = body.replace(r'\end{equation}', r'\qquad (1)'+'\n'+r'\end{equation}')
    for env, caption_prefix in [('figure', '図'), ('table', '表')]:
        count = 0
        pattern = r'\\begin\{'+env+r'\*?\}(?:\[[^\]]*\])?(.*?)\\end\{'+env+r'\*?\}'
        def convert(match):
            nonlocal count
            count += 1
            contents = match[1]
            caption = re.search(r'\\caption\{([^}]+)\}', contents)[1]
            english = re.search(r'\\ecaption\{([^}]+)\}', contents)[1]
            contents = re.sub(r'\\e?caption\{[^}]+\}', '', contents).replace(r'\centering', '')
            contents = contents.replace('.pdf}', '.png}')
            heading = r'\paragraph*{'+caption_prefix+' '+str(count)+'　'+caption+'}\n'
            english_line = r'\noindent '+english+'\n'
            return (contents+'\n'+heading+english_line if env == 'figure' else heading+english_line+contents)+'\n'
        body = re.sub(pattern, convert, body, flags=re.S)
        assert count == 2, (env, count)
    body = body.replace(r'\input{references.tex}', r'\section*{参考文献}'+'\n'+r'\begin{enumerate}'+'\n'+''.join(r'\item '+t.strip()+'\n\n' for _,t in refs)+r'\end{enumerate}')
    abstracts = read('abstracts.tex')
    parts = {env:re.search(r'\\begin\{'+env+r'\}(.*?)\\end\{'+env+r'\}',abstracts,re.S)[1].strip() for env in ['abstract','keyword','eabstract','ekeyword']}
    title = re.search(r'\\jtitle\{([^}]+)\}',main)[1]
    etitle = re.search(r'\\etitle\{([^}]+)\}',main)[1]
    preamble = r'\documentclass{article}'+'\n'+r'\usepackage{graphicx,amsmath,url}'+'\n'+r'\title{'+title+'}\n'+r'\author{著者未設定 / Author details pending}'+'\n'+r'\date{}'+'\n'+r'\begin{document}\maketitle'+'\n'+etitle+'\n\n'+r'\section*{あらまし}'+'\n'+parts['abstract']+'\n\n'+r'\noindent キーワード：'+parts['keyword']+'\n\n'
    ending = '\n'+r'\section*{Abstract}'+'\n'+parts['eabstract']+'\n\n'+r'\noindent Key words: '+parts['ekeyword']+'\n'+r'\end{document}'+'\n'
    return preamble + body + ending, numbers


def font_settings(style, east_asia='Yu Mincho'):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
    fonts.set(qn('w:eastAsia'), east_asia)


def build():
    pandoc = os.environ.get('PANDOC') or shutil.which('pandoc')
    if not pandoc:
        raise SystemExit('Set PANDOC to an installed Pandoc executable; see README.md.')
    build_dir = ROOT/'build'/'word'; build_dir.mkdir(parents=True, exist_ok=True)
    reference = Document()
    for section in reference.sections:
        section.page_width=Mm(210); section.page_height=Mm(297)
        section.top_margin=Mm(20); section.bottom_margin=Mm(20)
        section.left_margin=Mm(20); section.right_margin=Mm(20)
    for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Caption']:
        font_settings(reference.styles[name], 'Yu Gothic' if name.startswith('Heading') else 'Yu Mincho')
    reference.styles['Normal'].paragraph_format.space_after = Pt(6)
    reference.styles['Normal'].paragraph_format.line_spacing = 1.15
    reference.styles['Title'].font.size=Pt(16)
    reference.styles['Heading 1'].font.size=Pt(13)
    reference.styles['Heading 2'].font.size=Pt(11)
    reference_path = build_dir/'word_reference.docx'; reference.save(reference_path)
    source, numbers = make_source()
    source_path = ROOT/'word_source.tex'; source_path.write_text(source,encoding='utf-8')
    output = ROOT/'manuscript.docx'
    command=[pandoc,str(source_path),'--from=latex','--to=docx','--standalone','--number-sections','--metadata=lang:ja-JP','--resource-path='+str(ROOT),'--reference-doc='+str(reference_path),'--output='+str(output)]
    proc=subprocess.run(command,cwd=ROOT,text=True,capture_output=True,check=True)
    (build_dir/'pandoc.stdout.log').write_text(proc.stdout)
    (build_dir/'pandoc.stderr.log').write_text(proc.stderr)
    doc=Document(output)
    for table in doc.tables:
        table.style='Table Grid'
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size=Pt(9)
    doc.core_properties.title=re.search(r'\\jtitle\{([^}]+)\}',read('main.tex'))[1]
    doc.core_properties.author='著者未設定 / Author details pending'
    doc.core_properties.subject='Editable Word companion; experimental evidence unchanged from v04.'
    doc.save(output)
    with zipfile.ZipFile(output) as archive:
        assert archive.testzip() is None
        xml=etree.fromstring(archive.read('word/document.xml'))
        media=[n for n in archive.namelist() if n.startswith('word/media/')]
    plain=''.join(xml.xpath('//w:t/text()',namespaces=NS))
    expected=[['課題','Single','Separate','Inline'],['写真操作','139,871','139,871','129,374'],['精算','211,110','211,110','211,110'],['返金','102,087','102,087','102,087'],['CSV処理','440,461','332,588','394,351'],['全4課題','893,529','785,656','836,922']]
    tables=[[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
    assert tables[0] == expected, tables[0]
    expected_termination=[['Interface','制御','論理入出力','モデル呼出し','API要求'],['文書取得必須','通常継続','130,439','15','104'],['','証拠付き終了','91,116','13','100'],['公開名アクセス','通常継続','151,664','14','103'],['','証拠付き終了','112,233','12','99']]
    assert tables[1] == expected_termination, tables[1]
    assert len(tables)==2 and len(media)==2
    assert len(xml.xpath('//m:oMathPara',namespaces=NS))==1
    for marker in ['12.07%','6.53%','30.15%','26.00%','Reflexion','Generative agents','LLMLingua-2','使用量不明','著者未設定']:
        assert marker in plain, marker
    assert len(numbers)==12
    validation={'format':'docx','purpose':'editable one-column companion, not a verified IEICE Word template','pandoc_version':subprocess.check_output([pandoc,'--version'],text=True).splitlines()[0], 'python_docx_version':__import__('docx').__version__, 'source_sha256':{n:sha(ROOT/n) for n in ['main.tex','abstracts.tex','tables.tex','references.tex','word_source.tex','build_word.py','figures/control_points.png','figures/separate_studies.png']}, 'docx_sha256':sha(output),'paragraphs':len(doc.paragraphs),'editable_tables':len(tables),'embedded_images':len(media),'native_omml_display_equations':len(xml.xpath('//m:oMathPara',namespaces=NS)),'references':len(numbers),'citation_numbers':numbers,'scientific_table_cells_match':True,'pandoc_warnings':proc.stderr.strip(),'zip_integrity':True,'office_full_render_verified':False,'limits':['Word/LibreOffice was not installed; OOXML structure and content are validated, but pagination in Microsoft Word is not certified.','Figures are the unchanged v04 versions pending user selection.']}
    (ROOT/'word_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:validation[k] for k in ['docx_sha256','editable_tables','embedded_images','native_omml_display_equations','references','pandoc_warnings']},ensure_ascii=False))


if __name__ == '__main__':
    build()
