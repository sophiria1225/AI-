#!/usr/bin/env python3
"""Rebuild the editable manuscript on the unchanged official IEICE D template.

Usage: python build_word.py --template /path/to/iss_lettertemplete.docx
Requires python-docx and lxml. The official template is obtained separately.
"""
from __future__ import annotations
import argparse,copy,hashlib,json,re,zipfile
from datetime import datetime,timezone
from pathlib import Path
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parent
TEMPLATE_SHA='1b5b574780cfb971897bea26d07b9dd84a6f523124664a5bf6b84ea3fee75709'
URL='https://www.ieice.org/jpn/shiori/jd/iss_lettertemplete.docx?download=1'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def run(template):
 assert sha(template)==TEMPLATE_SHA,'Unexpected official template bytes'
 original=Document(template);doc=Document(template)
 base=copy.deepcopy(doc.element.body.sectPr)
 for child in list(doc.element.body):
  if child is not doc.element.body.sectPr:doc.element.body.remove(child)
 content=json.loads((ROOT/'word_content.json').read_text());data=content['items']
 def p(text='',style='共通_本文',source=None):
  paragraph=doc.add_paragraph(style=style)
  if source is None:paragraph.add_run(text)
  else:
   for xml in source['inline_xml']:paragraph._p.append(etree.fromstring(xml.encode()))
  return paragraph
 sections=0
 def boundary(columns,start='continuous'):
  nonlocal sections
  paragraph=p();paragraph.paragraph_format.space_before=Pt(0);paragraph.paragraph_format.space_after=Pt(0)
  paragraph.paragraph_format.line_spacing=Pt(1)
  pp=paragraph._p.get_or_add_pPr();snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
  mark=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'2');mark.append(sz);pp.append(mark)
  sect=copy.deepcopy(base);sect.find(qn('w:cols')).set(qn('w:num'),str(columns))
  sect.find(qn('w:type')).set(qn('w:val'),start)
  if sections:
   title=sect.find(qn('w:titlePg'))
   if title is not None:sect.remove(title)
  paragraph._p.get_or_add_pPr().append(sect);sections+=1
 def table(item,index):
  t=doc.add_table(rows=len(item['rows']),cols=len(item['rows'][0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  widths=([17,16,16,16] if index==1 else [12,13,17,13,10])
  for col,width in zip(t.columns,widths):col.width=Mm(width)
  for ri,(row,values) in enumerate(zip(t.rows,item['rows'])):
   for ci,(cell,value) in enumerate(zip(row.cells,values)):
    cell.width=Mm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;para=cell.paragraphs[0];para.style=doc.styles['共通_文献'];para.text=value;para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.first_line_indent=Mm(0);para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(0)
    ind=para._p.get_or_add_pPr().find(qn('w:ind'));ind.set(qn('w:firstLineChars'),'0');ind.set(qn('w:left'),'0');ind.set(qn('w:right'),'0');ind.set(qn('w:hanging'),'0')
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');para._p.get_or_add_pPr().append(snap)
    margins=OxmlElement('w:tcMar')
    for side in ['left','right']:
     edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'35');edge.set(qn('w:type'),'dxa');margins.append(edge)
    cell._tc.get_or_add_tcPr().append(margins)
    if ri>0 and ((index==1 and ci>=1) or (index==2 and ci>=2)):
     cell._tc.get_or_add_tcPr().append(OxmlElement('w:noWrap'))
    para.paragraph_format.keep_together=True
    para.paragraph_format.keep_with_next=True
    if index==2 and ri==2:para.paragraph_format.space_after=Pt(2)
    if ci>=2 or (index==1 and ci>=1):para.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    if ri==0:
     para.alignment=WD_ALIGN_PARAGRAPH.CENTER
     for r in para.runs:r.bold=True
   trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
   if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  pr=t._tbl.tblPr
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','bottom','left','right','insideH','insideV']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'4');borders.append(e)
  pr.append(borders)
  for cell in t.rows[0].cells:
   b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');b.append(e);cell._tc.get_or_add_tcPr().append(b)
  if index==1:
   for cell in t.rows[-1].cells:
    b=OxmlElement('w:tcBorders');edge=OxmlElement('w:top');edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');b.append(edge);cell._tc.get_or_add_tcPr().append(b)
  return t
 title=data[0]['text'];mark=content['title_line_break_before'];assert title.endswith(mark)
 p(title[:-len(mark)]+'\n'+mark,'研究速報_タイトル（和）')
 p('著者未設定','研究速報_著者名（和）')
 p(data[2]['text'],'研究速報_タイトル（英）')
 p('Author details pending','研究速報_著者名（英）')
 p('所属・会員種別・連絡先：未設定','研究速報_所属')
 p('あらまし　'+data[4]['text'])
 p('キーワード　'+data[5]['text'].split('：',1)[1])
 references=False;english_section=False;refnum=0;table_count=0;figure_count=0;after_heading=False;after_table_note=False;i=6
 while i<len(data):
  item=data[i]
  if item['kind']=='figure':
   figure_count+=1
   span=item['span'];assert span in ('page_top','column')
   if span=='page_top':boundary(2)
   image=p(style='共通_図');image.alignment=WD_ALIGN_PARAGRAPH.CENTER
   image.paragraph_format.space_before=Pt(0);image.paragraph_format.keep_with_next=True
   image.add_run().add_picture(str(ROOT/item['path']),width=Mm(item['width_mm']))
   for offset in (1,2):
    caption=p(data[i+offset]['text'],'共通_キャプション');caption.paragraph_format.keep_together=True
    caption.paragraph_format.keep_with_next=(offset==1)
   # The one-column figure section starts on a new page; the following body
   # section is continuous two-column text. Small figures stay inside a column.
   if span=='page_top':boundary(1,start='nextPage')
   i+=3;after_heading=False;continue
  if item['kind']=='equation':
   formula=p(style='共通_本文');fmt=formula.paragraph_format
   fmt.first_line_indent=Mm(0);fmt.left_indent=Mm(0);fmt.right_indent=Mm(0)
   fmt.tab_stops.clear_all();fmt.tab_stops.add_tab_stop(Mm(32),WD_TAB_ALIGNMENT.CENTER);fmt.tab_stops.add_tab_stop(Mm(71),WD_TAB_ALIGNMENT.RIGHT)
   ind=formula._p.get_or_add_pPr().find(qn('w:ind'));ind.set(qn('w:firstLineChars'),'0')
   formula.add_run('\t');formula._p.append(etree.fromstring(item['math_xml'].encode()));formula.add_run('\t('+str(item['number'])+')')
   fmt.keep_together=True;fmt.keep_with_next=True;i+=1;continue
  if item['kind']=='table':table_count+=1;table(item,table_count);i+=1;continue
  text=item['text'];sty=item['style'];role=item.get('role')
  if text=='参考文献':references=True;p('文　　　献','共通_文献見出し')
  elif text=='Abstract':
   references=False;english_section=True
   english=p('Abstract',style='Normal');english.paragraph_format.page_break_before=True
  elif references:
   refnum+=1;q=p(style='共通_文献');q.add_run('\t['+str(refnum)+']\t')
   for xml in item['inline_xml']:q._p.append(etree.fromstring(xml.encode()))
  elif sty in ('Heading 1','Heading 2'):
   heading=p(text.replace('\t','．' if sty=='Heading 1' else '　',1),'共通_節見出し' if sty=='Heading 1' else '共通_項見出し')
   heading.paragraph_format.keep_with_next=True;heading.paragraph_format.keep_together=True;after_heading=True
  elif role in ('table_caption','table_note'):
   caption=p(text,'共通_キャプション');caption.paragraph_format.keep_together=True
   if role=='table_note' and data[i-1]['kind']=='table':caption.paragraph_format.space_before=Pt(2)
   after_table_note=(role=='table_note' and not (i+1<len(data) and data[i+1].get('role')=='table_note'))
   caption.paragraph_format.keep_with_next=(role=='table_caption' or (i+1<len(data) and data[i+1].get('role')=='table_note'))
  elif english_section:
   english=p(text,style='Normal');english.alignment=WD_ALIGN_PARAGRAPH.LEFT
  else:
   body=p(source=item)
   body.paragraph_format.widow_control=True
   if (after_heading and len(text)<=150) or after_table_note or item.get('keep_together'):body.paragraph_format.keep_together=True
   if item.get('keep_with_next'):body.paragraph_format.keep_with_next=True
   after_heading=False;after_table_note=False
  i+=1
 final=doc.element.body.sectPr
 final.clear()
 for attr,val in base.attrib.items():final.set(attr,val)
 for child in base:final.append(copy.deepcopy(child))
 if sections:
  title=final.find(qn('w:titlePg'))
  if title is not None:final.remove(title)
 doc.core_properties.title=data[0]['text'];doc.core_properties.author='著者未設定'
 doc.core_properties.subject='IEICE Japanese D letter; official template based editable manuscript'
 doc.core_properties.last_modified_by='';doc.core_properties.comments=''
 stamp=datetime(2026,10,28,22,30,0,tzinfo=timezone.utc)
 doc.core_properties.created=stamp;doc.core_properties.modified=stamp
 for element in list(doc.core_properties._element):
  if etree.QName(element).localname=='lastPrinted':doc.core_properties._element.remove(element)
 target=ROOT/'manuscript.docx';doc.save(target)
 # Exact official package parts are retained byte-for-byte, not merely recreated styles.
 with zipfile.ZipFile(template) as z:preserved={n:z.read(n) for n in z.namelist() if n in ['word/styles.xml','word/fontTable.xml','word/settings.xml'] or re.match(r'word/(?:header|footer)\d+\.xml$',n)}
 with zipfile.ZipFile(target) as z:allparts={n:z.read(n) for n in z.namelist()}
 allparts.update(preserved)
 # Restore the sole v05 note using the unchanged official reference-text style.
 footnotes=etree.fromstring(allparts['word/footnotes.xml'])
 note=OxmlElement('w:footnote');note.set(qn('w:id'),'20')
 np=OxmlElement('w:p');pr=OxmlElement('w:pPr');sty=OxmlElement('w:pStyle');sty.set(qn('w:val'),'af2');pr.append(sty);np.append(pr)
 nr=OxmlElement('w:r');nr.append(OxmlElement('w:footnoteRef'));np.append(nr)
 nr=OxmlElement('w:r');nt=OxmlElement('w:t');nt.text=' https://github.com/sophiria1225/AI-';nt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');nr.append(nt);np.append(nr);note.append(np);footnotes.append(note)
 allparts['word/footnotes.xml']=etree.tostring(footnotes,xml_declaration=True,encoding='UTF-8',standalone=True)
 # Remove the template's two explanatory screenshots and their unused relations.
 relname='word/_rels/document.xml.rels'
 rels=etree.fromstring(allparts[relname])
 bodyxml=etree.fromstring(allparts['word/document.xml'])
 used=set(bodyxml.xpath('//@r:embed | //@r:link',namespaces={'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}))
 for rel in list(rels):
  if rel.get('Type','').endswith('/image') and rel.get('Id') not in used:
   allparts.pop('word/'+rel.get('Target'),None);rels.remove(rel)
 allparts[relname]=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
  for n,content in allparts.items():z.writestr(n,content)
 with zipfile.ZipFile(target) as z:
  assert z.testzip() is None
  xml=etree.fromstring(z.read('word/document.xml'));plain=''.join(xml.xpath('//w:t/text()',namespaces=NS))
  assert all(z.read(n)==content for n,content in preserved.items())
  for name in ['電子　花子','情報　太郎','通信　次郎','hanako@','Denshi','JOHO','TSUSHIN']:
   assert name not in plain,name
  sects=xml.xpath('//w:sectPr',namespaces=NS)
  for sect in sects:
   assert dict(sect.find(qn('w:pgSz')).attrib)==dict(base.find(qn('w:pgSz')).attrib)
   assert dict(sect.find(qn('w:pgMar')).attrib)==dict(base.find(qn('w:pgMar')).attrib)
   assert dict(sect.find(qn('w:docGrid')).attrib)==dict(base.find(qn('w:docGrid')).attrib)
  assert len(xml.xpath('//w:tbl',namespaces=NS))==2
  assert len(xml.xpath('//w:drawing',namespaces=NS))==3
  assert len([n for n in z.namelist() if n.startswith('word/media/')])==3
  assert sum(x['kind']=='equation' for x in data)==1
  assert len(xml.xpath('//m:nary',namespaces=NS))==1
  assert refnum==12
  notes=etree.fromstring(z.read('word/footnotes.xml'))
  assert 'https://github.com/sophiria1225/AI-' in ''.join(notes.xpath('//w:t/text()',namespaces=NS))
  assert {r.get(qn('w:id')) for r in xml.xpath('//w:footnoteReference',namespaces=NS)}=={'20'}
  assert len(notes.xpath('//w:footnote[@w:id="20"]',namespaces=NS))==1
  reread=Document(target)
  assert [[[c.text for c in r.cells] for r in t.rows] for t in reread.tables]==[x['rows'] for x in data if x['kind']=='table']
  # Verify every substantive source paragraph remains text-identical, ignoring heading numbering and field labels.
  canonical=json.loads((ROOT/'word_content.json').read_text())['items']
  for idx,item in enumerate(canonical):
   if item['kind']!='paragraph' or item['style'].startswith('Heading ') or idx in (1,5):continue
   assert item['text'] in plain,(idx,item['text'])
  assert '12.07%' in plain and '6.53%' in plain and '30.15%' in plain and '26.00%' in plain
  ja=len(canonical[4]['text'])
  abstract_index=next(n for n,x in enumerate(canonical) if x['kind']=='paragraph' and x['text']=='Abstract')
  en=len(re.findall(r"[A-Za-z0-9]+(?:['’][A-Za-z]+)?",canonical[abstract_index+1]['text']))
  assert ja<=120 and en<=50
  validation={'format':'docx','purpose':'manuscript rebuilt on official IEICE Japanese D letter template','official_template_url':URL,'official_template_sha256':TEMPLATE_SHA,'docx_sha256':sha(target),'source_sha256':{n:sha(ROOT/n) for n in ['build_word.py','word_content.json','main.tex','abstracts.tex','tables.tex','references.tex','figures/control_points.png','figures/result_summary.png','figures/result_termination.png']},'official_parts_byte_identical':{n:hashlib.sha256(b).hexdigest() for n,b in preserved.items()},'page_geometry_twips':{'width':10319,'height':14572,'left':851,'right':851,'top':1418,'bottom':1418,'column_space':443,'grid_line_pitch':266},'section_column_counts':[int(s.find(qn('w:cols')).get(qn('w:num'),'1')) for s in sects],'body_style':'共通_本文','body_size_pt':8.5,'editable_tables':2,'embedded_manuscript_images':3,'figure_widths_mm':[130,65,65],'native_omml_display_equations':1,'equation_number_layout':'native OMML with independent right-aligned tab number','references':12,'japanese_abstract_chars':ja,'english_abstract_word_count_conservative':en,'author_details':'pending','scientific_table_numbers_match_v08':True,'source_body_paragraphs_preserved':True,'repository_url_footnote_preserved':True,'table_width_mm':65,'table_font_size_pt':7,'table_line_spacing_pt':10.8,'english_abstract_starts_new_page':True,'english_abstract_columns':2,'full_render_verified':False,'limits':['Pagination and layout require separate full-page rendering validation.','MS Mincho is retained in DOCX; a local PDF proof may substitute Hiragino Mincho ProN.','Publication year, volume and page placeholders are retained from the official template; author details remain unset.']}
  (ROOT/'word_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
  provenance={'official_template':{'url':URL,'sha256':TEMPLATE_SHA,'bytes':Path(template).stat().st_size},'basis':'Official Japanese D letter Word template; existing styles and header/footer parts unchanged','content_source_version':'v09','template_redistributed':False,'python_docx_version':__import__('docx').__version__,'rebuild_command':'python build_word.py --template /path/to/iss_lettertemplete.docx','body_geometry':'Actual template XML (B5), not explanatory prose in template','figures':'Figure 1 is a two-panel control diagram at the top of a page; figures 2 and 3 are independent single-column result plots. Prior alternatives remain in v08.'}
  (ROOT/'official_template_provenance.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps({'docx':str(target),'sha256':sha(target),'sections':len(sects),'refs':refnum},ensure_ascii=False))
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--template',type=Path,required=True);args=ap.parse_args();run(args.template)
